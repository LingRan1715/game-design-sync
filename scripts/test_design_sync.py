"""Behavioral regression tests in isolated workspaces; no real UE installation needed."""
import argparse
import base64
import copy
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import design_sync as sync
import impact_scan
import ue_bridge
import ue_editor_bridge


STARTER = Path(__file__).resolve().parents[1] / "assets" / "starter"


class SyncTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="gds-test-")
        self.root = Path(self.temp.name) / "project"
        shutil.copytree(STARTER, self.root)
        self.manifest = self.root / "design-sync.json"
        self.init()

    def tearDown(self):
        self.temp.cleanup()

    def call(self, command, migration=False):
        return sync.execute(argparse.Namespace(command=command, manifest=str(self.manifest), accept_reviewed_migration=migration))

    def init(self):
        self.assertEqual(self.call("init")[1], 0)

    def json_edit(self, **values):
        p = self.root / "GameData/skills.json"
        rows = sync.load_json(p.read_bytes())
        rows[0].update(values)
        p.write_bytes(sync.encode(rows))

    def csv_edit(self, before, after):
        p = self.root / "Tables/skills.csv"
        p.write_text(p.read_text(encoding="utf-8-sig").replace(before, after), encoding="utf-8")

    def md_edit(self, before, after):
        p = self.root / "Design/skills.md"
        p.write_text(p.read_text(encoding="utf-8").replace(before, after), encoding="utf-8")

    def data(self):
        root, cfg, _ = sync.config(self.manifest)
        ds = cfg["datasets"][0]
        return [sync.View(root, ds, v).records for v in ds["views"]]

    def save_config(self, cfg):
        self.manifest.write_bytes(sync.encode(cfg))

    def configure(self, mutate):
        cfg = sync.load_json(self.manifest.read_bytes())
        mutate(cfg)
        self.save_config(cfg)
        self.call("rebaseline", True)

    def snapshot(self):
        return {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file() and ".design-sync" not in p.parts}

    def test_design_to_table_and_runtime(self):
        self.md_edit('"damage": 100', '"damage": 120')
        self.assertEqual(self.call("check")[1], 1)
        self.assertEqual(self.call("plan")[0]["status"], "drift")
        self.assertEqual(self.call("apply")[1], 0)
        self.assertTrue(all(v["fireball"]["damage"] == 120 for v in self.data()))
        self.assertIn("示例备注不参与跨文件同步", (self.root / "Tables/skills.csv").read_text(encoding="utf-8-sig"))
        self.assertEqual(self.call("check")[1], 0)

    def test_runtime_to_design_and_table(self):
        self.json_edit(Cooldown=3)
        self.call("apply")
        self.assertTrue(all(v["fireball"]["cooldown"] == 3 for v in self.data()))

    def test_excel_csv_to_design_and_runtime(self):
        self.csv_edit("fireball,100,2.0,true", "fireball,125,2.0,true")
        self.call("apply")
        self.assertTrue(all(v["fireball"]["damage"] == 125 for v in self.data()))

    def test_independent_changes_merge(self):
        self.md_edit('"damage": 100', '"damage": 120')
        self.json_edit(Cooldown=3)
        self.call("apply")
        self.assertTrue(all(v["fireball"]["damage"] == 120 and v["fireball"]["cooldown"] == 3 for v in self.data()))

    def test_convergent_changes_are_not_conflicts(self):
        self.md_edit('"damage": 100', '"damage": 120')
        self.json_edit(Damage=120)
        self.assertEqual(self.call("apply")[1], 0)

    def test_conflict_never_writes(self):
        self.md_edit('"damage": 100', '"damage": 120')
        self.json_edit(Damage=140)
        before = self.snapshot()
        state = (self.root / ".design-sync/state.json").read_bytes()
        result, code = self.call("apply")
        self.assertEqual(code, 2)
        self.assertEqual(result["conflicts"][0]["base"], 100)
        self.assertEqual(before, self.snapshot())
        self.assertEqual(state, (self.root / ".design-sync/state.json").read_bytes())

    def test_delete_versus_edit_conflict(self):
        (self.root / "GameData/skills.json").write_text("[]", encoding="utf-8")
        self.md_edit('"damage": 100', '"damage": 120')
        report, code = self.call("apply")
        self.assertEqual(code, 2)
        self.assertEqual(report["conflicts"][0]["reason"], "delete-versus-edit")

    def test_delete_propagates_without_checks_that_require_rows(self):
        self.configure(lambda cfg: cfg.update(checks=[]))
        (self.root / "GameData/skills.json").write_text("[]", encoding="utf-8")
        self.call("apply")
        self.assertTrue(all(v == {} for v in self.data()))

    def test_new_complete_row_propagates(self):
        p = self.root / "GameData/skills.json"
        rows = sync.load_json(p.read_bytes())
        rows.append({"Name": "icebolt", "Damage": 60, "Cooldown": 1, "Enabled": False})
        p.write_bytes(sync.encode(rows))
        self.call("apply")
        self.assertTrue(all(v["icebolt"]["enabled"] is False for v in self.data()))

    def test_schema_range_and_strict_type_fail_before_writes(self):
        for value in (-1, True, 1.2, "200"):
            with self.subTest(value=value):
                self.json_edit(Damage=value)
                before = self.snapshot()
                with self.assertRaises(sync.SyncError):
                    self.call("apply")
                self.assertEqual(before, self.snapshot())

    def test_duplicate_ids_and_case_collision(self):
        p = self.root / "GameData/skills.json"
        rows = sync.load_json(p.read_bytes())
        rows.append(copy.deepcopy(rows[0]))
        p.write_bytes(sync.encode(rows))
        with self.assertRaisesRegex(sync.SyncError, "duplicate"):
            self.call("plan")
        ds = {"id": "x", "id_field": "id", "fields": {"id": {"type": "string"}}, "case_insensitive_ids": True}
        p.write_bytes(sync.encode([{"id": "Fireball"}, {"id": "fireball"}]))
        with self.assertRaisesRegex(sync.SyncError, "collision"):
            sync.View(self.root, ds, {"path": "GameData/skills.json", "format": "json"})

    def test_duplicate_json_keys_rejected(self):
        with self.assertRaisesRegex(sync.SyncError, "Duplicate JSON key"):
            sync.load_json(b'{"a":1,"a":2}')

    def test_failed_check_restores_all_files_and_baseline(self):
        asset = self.root / "dummy.uasset"
        asset.write_bytes(b"original-asset")
        self.configure(lambda cfg: cfg.update(rollback_paths=["dummy.uasset"], checks=[{
            "id": "deliberate-failure", "argv": ["{python}", "-c", "from pathlib import Path; Path('dummy.uasset').write_bytes(b'changed'); raise SystemExit(7)"], "timeout": 10}]))
        self.md_edit('"damage": 100', '"damage": 120')
        before, state = self.snapshot(), (self.root / ".design-sync/state.json").read_bytes()
        with self.assertRaisesRegex(sync.SyncError, "Check failed"):
            self.call("apply")
        self.assertEqual(before, self.snapshot())
        self.assertEqual(state, (self.root / ".design-sync/state.json").read_bytes())
        self.assertFalse((self.root / ".design-sync/pending.json").exists())

    def test_timeout_rolls_back(self):
        self.configure(lambda cfg: cfg.update(checks=[{"id": "slow", "argv": ["{python}", "-c", "import time; time.sleep(5)"], "timeout": 0.1}]))
        self.json_edit(Damage=120)
        before = self.snapshot()
        with self.assertRaisesRegex(sync.SyncError, "timed out"):
            self.call("apply")
        self.assertEqual(before, self.snapshot())

    def test_validator_cannot_change_synced_data(self):
        self.configure(lambda cfg: cfg.update(checks=[{"id": "mutator", "argv": ["{python}", "-c", "from pathlib import Path; Path('GameData/skills.json').write_text('[]')"]}]))
        self.json_edit(Damage=120)
        before = self.snapshot()
        with self.assertRaisesRegex(sync.SyncError, "changed synchronized data"):
            self.call("apply")
        self.assertEqual(before, self.snapshot())

    def test_idempotence_preserves_source_bytes(self):
        before = self.snapshot()
        self.call("apply")
        self.call("apply")
        self.assertEqual(before, self.snapshot())

    def test_mapping_changes_need_reviewed_migration(self):
        cfg = sync.load_json(self.manifest.read_bytes())
        cfg["datasets"][0]["fields"]["damage"]["max"] = 9999
        self.save_config(cfg)
        with self.assertRaisesRegex(sync.SyncError, "Manifest changed"):
            self.call("check")
        with self.assertRaisesRegex(sync.SyncError, "requires"):
            self.call("rebaseline")
        self.call("rebaseline", True)
        self.assertEqual(self.call("check")[1], 0)

    def test_init_does_not_erase_existing_drift(self):
        self.json_edit(Damage=120)
        with self.assertRaisesRegex(sync.SyncError, "Baseline exists"):
            self.call("init")

    def test_first_init_refuses_disagreement(self):
        (self.root / ".design-sync/state.json").unlink()
        self.json_edit(Damage=120)
        with self.assertRaisesRegex(sync.SyncError, "disagree"):
            self.call("init")

    def test_path_escape_rejected(self):
        cfg = sync.load_json(self.manifest.read_bytes())
        cfg["datasets"][0]["views"][0]["path"] = "../outside.md"
        self.save_config(cfg)
        with self.assertRaisesRegex(sync.SyncError, "escapes"):
            self.call("plan")

    def test_pending_recovery(self):
        paths = [self.root / "GameData/skills.json"]
        before = sync.snapshots(self.root, paths)
        state = (self.root / ".design-sync/state.json").read_bytes()
        pending = {"before": before, "state_before": base64.b64encode(state).decode("ascii")}
        (self.root / ".design-sync/pending.json").write_bytes(sync.encode(pending))
        self.json_edit(Damage=222)
        with self.assertRaisesRegex(sync.SyncError, "Interrupted"):
            self.call("plan")
        self.assertEqual(self.call("recover")[0]["status"], "recovered")
        self.assertEqual(sync.snapshots(self.root, paths), before)

    def test_recovery_does_not_depend_on_valid_current_manifest(self):
        paths = [self.manifest]
        before = sync.snapshots(self.root, paths)
        state = (self.root / ".design-sync/state.json").read_bytes()
        pending = {"before": before, "state_before": base64.b64encode(state).decode("ascii")}
        (self.root / ".design-sync/pending.json").write_bytes(sync.encode(pending))
        self.manifest.write_text("invalid-json")
        self.assertEqual(self.call("recover")[0]["status"], "recovered")
        self.assertEqual(sync.snapshots(self.root, paths), before)

    def test_check_mutation_of_manifest_is_restored(self):
        self.configure(lambda cfg: cfg.update(checks=[{"id": "bad-config-writer", "argv": ["{python}", "-c", "from pathlib import Path; Path('design-sync.json').write_text('{}')"]}]))
        before = self.manifest.read_bytes()
        self.json_edit(Damage=120)
        with self.assertRaisesRegex(sync.SyncError, "changed the manifest"):
            self.call("apply")
        self.assertEqual(before, self.manifest.read_bytes())

    def test_cross_dataset_foreign_key(self):
        cfg = sync.load_json(self.manifest.read_bytes())
        ds = cfg["datasets"][0]
        ds["fields"]["target"] = {"type": "string", "ref": "skills"}
        rows = {"fireball": {"id": "fireball", "damage": 100, "cooldown": 2.0, "enabled": True, "target": "missing"}}
        with self.assertRaisesRegex(sync.SyncError, "Broken reference"):
            sync.validate_refs(cfg, {"skills": rows})

    def test_json_path_and_unmapped_metadata_preserved(self):
        p = self.root / "GameData/skills.json"
        rows = sync.load_json(p.read_bytes())
        rows[0]["DeveloperNote"] = "preserve me"
        p.write_bytes(sync.encode({"meta": {"version": 42}, "rows": rows}))
        self.configure(lambda cfg: cfg["datasets"][0]["views"][2].update(json_path=["rows"]))
        self.md_edit('"damage": 100', '"damage": 120')
        # The example checker expects a root list; replace with no checks for this adapter test.
        cfg = sync.load_json(self.manifest.read_bytes())
        cfg["checks"] = []
        self.save_config(cfg)
        # Reconcile migration without dropping the local design edit.
        self.md_edit('"damage": 120', '"damage": 100')
        self.call("rebaseline", True)
        self.md_edit('"damage": 100', '"damage": 120')
        self.call("apply")
        result = sync.load_json(p.read_bytes())
        self.assertEqual(result["meta"]["version"], 42)
        self.assertEqual(result["rows"][0]["DeveloperNote"], "preserve me")
        self.assertEqual(result["rows"][0]["Damage"], 120)

    def test_semantic_code_change_and_removed_mapping_detected(self):
        manifest = self.root / "design-links.json"
        impact_scan.scan(manifest, True)
        p = self.root / "Tests/check_runtime_data.py"
        p.write_text(p.read_text(encoding="utf-8") + "\n# behavior review needed\n", encoding="utf-8")
        report, code = impact_scan.scan(manifest)
        self.assertEqual(code, 1)
        self.assertEqual(report["features"][0]["id"], "skill.fireball")
        cfg = sync.load_json(manifest.read_bytes())
        cfg["features"][0]["files"].remove("Tests/check_runtime_data.py")
        manifest.write_bytes(sync.encode(cfg))
        self.assertTrue(impact_scan.scan(manifest)[0]["mapping_changed"])

    def test_missing_semantic_file_cannot_be_acknowledged(self):
        (self.root / "Tests/check_runtime_data.py").unlink()
        with self.assertRaisesRegex(sync.SyncError, "Missing mapped file"):
            impact_scan.scan(self.root / "design-links.json", True)

    def test_cli_real_process_exit_codes(self):
        script = Path(sync.__file__)
        self.json_edit(Damage=120)
        result = subprocess.run([sys.executable, str(script), "check", "--manifest", str(self.manifest)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(result.stdout)["status"], "drift")

    def test_xlsx_bidirectional_and_formula_refusal(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("optional openpyxl not installed")
        path = self.root / "Tables/skills.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Skills"
        sheet.append(["id", "damage", "cooldown", "enabled", "note"])
        sheet.append(["fireball", 100, 2.0, True, "keep"])
        book.create_sheet("Notes")["A1"] = "other sheet stays"
        book.save(path)
        self.configure(lambda cfg: cfg["datasets"][0]["views"].append({"id": "excel", "path": "Tables/skills.xlsx", "format": "xlsx", "sheet": "Skills", "data_only_sheet": True}))
        book = openpyxl.load_workbook(path)
        book["Skills"]["B2"] = 150
        book.save(path)
        self.call("apply")
        self.assertTrue(all(v["fireball"]["damage"] == 150 for v in self.data()))
        self.json_edit(Cooldown=4)
        self.call("apply")
        book = openpyxl.load_workbook(path)
        self.assertEqual(book["Skills"]["C2"].value, 4)
        self.assertEqual(book["Skills"]["E2"].value, "keep")
        self.assertEqual(book["Notes"]["A1"].value, "other sheet stays")
        book["Skills"]["B2"] = "=100+20"
        book.save(path)
        with self.assertRaisesRegex(sync.SyncError, "formulas"):
            self.call("apply")

    def bridge_setup(self):
        project = self.root / "Sample.uproject"
        project.write_text("{}")
        asset = self.root / "Content/Data/DT_Skills.uasset"
        asset.parent.mkdir(parents=True)
        asset.write_bytes(b"original")
        args = argparse.Namespace(command="pull", project=str(project), engine=sys.executable,
                                  bindings=str(self.root / "ue-bindings.example.json"), timeout=10)
        return args, asset

    def fake_ue(self, **changes):
        def run(argv, **kwargs):
            request_path = Path(kwargs["env"]["GDS_UE_REQUEST"])
            request = sync.load_json(request_path.read_bytes())
            for table in request["tables"]:
                rows = sync.load_json(Path(table["source"]).read_bytes())
                rows[0].update(changes)
                Path(table["staged"]).write_bytes(sync.encode(rows))
            Path(str(request_path) + ".result.json").write_bytes(sync.encode({"token": request["token"], "success": True, "operations": [{"rows": 1}]}))
            return subprocess.CompletedProcess(argv, 0, "fake UE adapter test only", "")
        return run

    def test_bridge_pull_refuses_local_json_edits(self):
        args, _ = self.bridge_setup()
        self.json_edit(Damage=120)
        original = (self.root / "GameData/skills.json").read_bytes()
        with patch.object(ue_bridge.subprocess, "run", side_effect=self.fake_ue(Damage=140)):
            with self.assertRaisesRegex(sync.SyncError, "overwrite"):
                ue_bridge.execute(args)
        self.assertEqual(original, (self.root / "GameData/skills.json").read_bytes())

    def test_bridge_pull_accepts_engine_changes_against_clean_baseline(self):
        args, _ = self.bridge_setup()
        with patch.object(ue_bridge.subprocess, "run", side_effect=self.fake_ue(Damage=140)):
            self.assertEqual(ue_bridge.execute(args)["status"], "pull-complete")
        self.call("apply")
        self.assertTrue(all(v["fireball"]["damage"] == 140 for v in self.data()))

    def test_bridge_failure_restores_assets(self):
        args, asset = self.bridge_setup()
        args.command = "push"
        def fail(argv, **kwargs):
            asset.write_bytes(b"partial-import")
            return subprocess.CompletedProcess(argv, 1, "failed", "")
        with patch.object(ue_bridge.subprocess, "run", side_effect=fail):
            with self.assertRaisesRegex(sync.SyncError, "UE operation failed"):
                ue_bridge.execute(args)
        self.assertEqual(asset.read_bytes(), b"original")

    def test_editor_bridge_numeric_tolerance_and_missing_fields(self):
        self.assertTrue(ue_editor_bridge.equal({"v": 0.1}, {"v": 0.10000000149}))
        self.assertFalse(ue_editor_bridge.equal({"Damage": 100}, {"Damage": 100, "Missing": 0}))
        self.assertFalse(ue_editor_bridge.equal(True, 1))


if __name__ == "__main__":
    unittest.main(verbosity=2)
