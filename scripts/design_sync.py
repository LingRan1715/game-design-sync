#!/usr/bin/env python3
"""Manifest-scoped, baseline-aware synchronization of design data. Python 3.10+."""
import argparse
import base64
import copy
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time


class SyncError(Exception):
    pass


def encode(value):
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")


def digest(data):
    return hashlib.sha256(data).hexdigest()


def atomic(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".gds-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def scoped(root, name):
    # Normalize the boundary too: Windows temp paths can contain 8.3 aliases,
    # and callers may pass an equivalent path containing '..' or a junction.
    root = root.resolve()
    if not isinstance(name, str) or not name or Path(name).is_absolute():
        raise SyncError(f"Expected project-relative path: {name!r}")
    path = (root / name).resolve()
    if not path.is_relative_to(root) or path == root:
        raise SyncError(f"Path escapes project: {name}")
    return path


def load_json(data):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise SyncError(f"Duplicate JSON key: {key}")
            result[key] = value
        return result
    return json.loads(data.decode("utf-8-sig"), object_pairs_hook=pairs,
                      parse_constant=lambda v: (_ for _ in ()).throw(SyncError(f"Invalid number: {v}")))


def config(path):
    path = path.resolve()
    root = path.parent
    raw = path.read_bytes()
    cfg = load_json(raw)
    if cfg.get("version") != 1 or not cfg.get("datasets"):
        raise SyncError("Manifest needs version=1 and nonempty datasets")
    seen_paths, seen_sets = set(), set()
    for ds in cfg["datasets"]:
        if not isinstance(ds.get("id"), str) or ds["id"] in seen_sets:
            raise SyncError("Dataset IDs must be unique strings")
        seen_sets.add(ds["id"])
        fields = ds.get("fields", {})
        if not fields or ds.get("id_field") not in fields:
            raise SyncError("Dataset needs fields and id_field")
        if fields[ds["id_field"]].get("type") != "string":
            raise SyncError("Stable ID must have string type")
        for spec in fields.values():
            if spec.get("type") not in ("string", "integer", "number", "boolean"):
                raise SyncError("Supported field types: string, integer, number, boolean")
            if spec.get("ref") and not isinstance(spec["ref"], str):
                raise SyncError("ref must be a dataset ID")
            if spec["type"] == "string" and spec.get("nullable"):
                raise SyncError("Nullable strings cannot round-trip through CSV; use explicit empty-string semantics")
        if len(ds.get("views", [])) < 2:
            raise SyncError("Each dataset needs at least two views")
        view_ids = set()
        for view in ds["views"]:
            if not isinstance(view.get("id"), str) or view["id"] in view_ids:
                raise SyncError("View IDs must be unique strings per dataset")
            view_ids.add(view["id"])
            p = scoped(root, view["path"])
            if p == path or ".design-sync" in p.relative_to(root).parts:
                raise SyncError("Managed files cannot be the manifest or state directory")
            if p in seen_paths:
                raise SyncError("One managed view per file; split docs/tables or use an adapter")
            seen_paths.add(p)
            if view.get("format") not in ("json", "csv", "markdown", "xlsx"):
                raise SyncError("Supported formats: json, csv, markdown, xlsx")
            aliases = view.get("columns", {})
            if set(aliases) - set(fields):
                raise SyncError("Unknown logical field in columns")
            names = [aliases.get(f, f) for f in fields]
            if len(set(names)) != len(names):
                raise SyncError("Column aliases must be distinct")
            if view["format"] == "xlsx" and (not view.get("sheet") or not view.get("data_only_sheet")):
                raise SyncError("XLSX requires sheet and data_only_sheet=true (no formulas/charts/macros)")
            if view["format"] == "xlsx" and p.suffix.lower() != ".xlsx":
                raise SyncError("Only .xlsx data workbooks are supported, no macro-enabled files")
    for ds in cfg["datasets"]:
        for spec in ds["fields"].values():
            if spec.get("ref") and spec["ref"] not in seen_sets:
                raise SyncError("ref points to unknown dataset")
    for name in cfg.get("rollback_paths", []):
        p = scoped(root, name)
        if p == path or ".design-sync" in p.relative_to(root).parts:
            raise SyncError("Invalid rollback path")
    for check in cfg.get("checks", []):
        if not isinstance(check.get("argv"), list) or not check["argv"] or not all(isinstance(a, str) for a in check["argv"]):
            raise SyncError("Checks require argv string arrays; shell commands are unsupported")
        if check.get("timeout", 120) <= 0:
            raise SyncError("Check timeout must be positive")
    return root, cfg, digest(raw)


def coerce(value, spec, textual=False):
    kind = spec["type"]
    if value is None or (textual and value == "" and kind != "string"):
        if spec.get("nullable"):
            return None
        raise SyncError("Non-null field is empty")
    try:
        if kind == "string":
            if not isinstance(value, str):
                raise ValueError("expected string; use text cells for IDs")
            result = value
        elif kind == "boolean":
            if isinstance(value, bool):
                result = value
            elif textual and str(value).lower() in ("true", "false"):
                result = str(value).lower() == "true"
            else:
                raise ValueError("expected true/false")
        elif kind == "integer":
            if isinstance(value, bool):
                raise ValueError("boolean is not integer")
            if isinstance(value, int):
                result = value
            elif textual and isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value):
                result = int(value)
            else:
                raise ValueError("expected integer")
        else:
            if isinstance(value, bool) or (not textual and not isinstance(value, (int, float))):
                raise ValueError("expected number")
            result = float(value)
            if not math.isfinite(result):
                raise ValueError("expected finite number")
        if "enum" in spec and result not in spec["enum"]:
            raise ValueError("not in enum")
        if kind in ("integer", "number"):
            if "min" in spec and result < spec["min"]:
                raise ValueError(f"below minimum {spec['min']}")
            if "max" in spec and result > spec["max"]:
                raise ValueError(f"above maximum {spec['max']}")
        if kind == "string" and "pattern" in spec and not re.fullmatch(spec["pattern"], result):
            raise ValueError("pattern mismatch")
        return result
    except (ValueError, TypeError, OverflowError) as exc:
        raise SyncError(f"Invalid {kind} {value!r}: {exc}") from exc


class View:
    def __init__(self, root, ds, cfg):
        self.ds, self.cfg = ds, cfg
        self.path = scoped(root, cfg["path"])
        self.raw = self.path.read_bytes()
        self.kind = cfg["format"]
        self.names = {f: cfg.get("columns", {}).get(f, f) for f in ds["fields"]}
        if self.kind == "json":
            self.document = load_json(self.raw)
            self.rows = self.document
            for key in cfg.get("json_path", []):
                self.rows = self.rows[key]
        elif self.kind == "markdown":
            self.text = self.raw.decode("utf-8-sig")
            marker = re.escape(cfg.get("block", ds["id"]))
            matches = list(re.finditer(r"(<!-- design-sync:" + marker + r":begin -->\s*```json\s*\r?\n)(.*?)(\r?\n```\s*<!-- design-sync:" + marker + r":end -->)", self.text, re.S))
            if len(matches) != 1:
                raise SyncError(f"Expected exactly one marked JSON block: {self.path}")
            self.match = matches[0]
            self.rows = load_json(self.match[2].encode("utf-8"))
        elif self.kind == "csv":
            reader = csv.DictReader(io.StringIO(self.raw.decode("utf-8-sig"), newline=""))
            self.headers = reader.fieldnames or []
            if len(self.headers) != len(set(self.headers)) or not self.headers:
                raise SyncError("CSV needs unique headers")
            self.rows = list(reader)
            if any(None in row or any(v is None for v in row.values()) for row in self.rows):
                raise SyncError("CSV row width differs from header")
        else:
            try:
                import openpyxl
            except ImportError as exc:
                raise SyncError("XLSX adapter needs openpyxl in this Python environment") from exc
            self.book = openpyxl.load_workbook(io.BytesIO(self.raw), data_only=False)
            self.sheet = self.book[cfg["sheet"]]
            # This is a data-sheet adapter; it never reorders formula or layout sheets.
            if self.sheet.merged_cells.ranges or self.sheet.tables or self.sheet._charts or self.sheet._images:
                raise SyncError("XLSX sheet must be a plain data sheet")
            values = list(self.sheet.values)
            self.headers = list(values[0]) if values else []
            if not self.headers or any(not isinstance(h, str) for h in self.headers) or len(set(self.headers)) != len(self.headers):
                raise SyncError("XLSX needs unique string headers on row 1")
            if any(cell.data_type == "f" for row in self.sheet for cell in row):
                raise SyncError("XLSX formulas are not writable sync fields; export evaluated data separately")
            self.rows = [dict(zip(self.headers, r)) for r in values[1:] if any(v is not None for v in r)]
        if not isinstance(self.rows, list) or any(not isinstance(row, dict) for row in self.rows):
            raise SyncError("View must contain a list of objects")
        if self.kind in ("csv", "xlsx") and set(self.names.values()) - set(self.headers):
            raise SyncError(f"Missing required columns: {self.path}")
        self.records, self.originals = {}, {}
        for row in self.rows:
            record = {}
            for f, spec in ds["fields"].items():
                if self.names[f] not in row:
                    raise SyncError(f"Missing field {self.names[f]}: {self.path}")
                try:
                    record[f] = coerce(row[self.names[f]], spec, self.kind == "csv")
                except SyncError as exc:
                    raise SyncError(f"{self.path}, field {f}: {exc}") from exc
            rid = record[ds["id_field"]]
            if not rid or rid in self.records:
                raise SyncError(f"Empty or duplicate stable ID {rid!r}: {self.path}")
            self.records[rid], self.originals[rid] = record, row
        if ds.get("case_insensitive_ids") and len({rid.casefold() for rid in self.records}) != len(self.records):
            raise SyncError("Case-insensitive row ID collision")

    def render(self, records):
        if records == self.records:
            return self.raw
        rows = []
        for rid in list(self.records) + sorted(set(records) - set(self.records)):
            if rid not in records:
                continue
            row = copy.deepcopy(self.originals.get(rid, {}))
            for f, value in records[rid].items():
                row[self.names[f]] = value
            rows.append(row)
        if self.kind == "json":
            keys = self.cfg.get("json_path", [])
            document = copy.deepcopy(self.document)
            if keys:
                parent = document
                for key in keys[:-1]:
                    parent = parent[key]
                parent[keys[-1]] = rows
            else:
                document = rows
            return encode(document)
        if self.kind == "markdown":
            return (self.text[:self.match.start(2)] + encode(rows).decode("utf-8").rstrip("\n") + self.text[self.match.end(2):]).encode("utf-8")
        if self.kind == "csv":
            stream = io.StringIO(newline="")
            writer = csv.DictWriter(stream, fieldnames=self.headers, lineterminator="\r\n")
            writer.writeheader()
            writer.writerows({k: (str(v).lower() if isinstance(v, bool) else "" if v is None else v) for k, v in row.items()} for row in rows)
            return stream.getvalue().encode("utf-8-sig")
        # Preserve cell styles in remaining rows and preserve other sheets.
        count = max(self.sheet.max_row - 1, len(rows))
        for i in range(count):
            for j, header in enumerate(self.headers, 1):
                value = rows[i].get(header) if i < len(rows) else None
                if isinstance(value, str) and value.startswith("="):
                    raise SyncError("XLSX values beginning with '=' need a dedicated text adapter")
                self.sheet.cell(i + 2, j).value = value
        stream = io.BytesIO()
        self.book.save(stream)
        return stream.getvalue()


def merge(ds, views, baseline):
    merged, conflicts, events = {}, [], []
    all_ids = set(baseline)
    for view in views:
        all_ids.update(view.records)
    for rid in sorted(all_ids):
        old = baseline.get(rid)
        entries = [(v.cfg["id"], v.records.get(rid)) for v in views]
        absent = [name for name, row in entries if row is None]
        present = [(name, row) for name, row in entries if row is not None]
        if old is not None and absent:
            if any(row != old for _, row in present):
                conflicts.append({"dataset": ds["id"], "id": rid, "field": "*", "reason": "delete-versus-edit", "base": old, "views": dict(entries)})
            else:
                events.append({"dataset": ds["id"], "id": rid, "operation": "delete", "sources": absent})
            continue
        if not present:
            continue
        row = {}
        for f in ds["fields"]:
            changes = [(name, value[f]) for name, value in present if old is None or value[f] != old[f]]
            unique = []
            for _, value in changes:
                if value not in unique:
                    unique.append(value)
            if len(unique) > 1:
                conflicts.append({"dataset": ds["id"], "id": rid, "field": f, "reason": "divergent-edits", "base": old[f] if old else None, "views": dict(changes)})
            else:
                row[f] = unique[0] if unique else old[f]
                if changes:
                    events.append({"dataset": ds["id"], "id": rid, "field": f, "base": old[f] if old else None, "value": row[f], "sources": [name for name, _ in changes]})
        if len(row) == len(ds["fields"]):
            merged[rid] = row
    return merged, conflicts, events


def validate_refs(cfg, datasets):
    for ds in cfg["datasets"]:
        records = datasets[ds["id"]]
        if ds.get("case_insensitive_ids") and len({rid.casefold() for rid in records}) != len(records):
            raise SyncError("Case-insensitive row ID collision after merge")
        for rid, row in records.items():
            for f, spec in ds["fields"].items():
                coerce(row[f], spec)
                if spec.get("ref") and row[f] is not None and row[f] not in datasets[spec["ref"]]:
                    raise SyncError(f"Broken reference {ds['id']}.{rid}.{f} -> {spec['ref']}:{row[f]}")


def inspect(root, cfg, state, signature):
    if state["manifest_hash"] != signature:
        raise SyncError("Manifest changed; migrate mappings/schema and use rebaseline after review")
    datasets, all_views, conflicts, events = {}, [], [], []
    for ds in cfg["datasets"]:
        views = [View(root, ds, v) for v in ds["views"]]
        result, problems, changes = merge(ds, views, state["datasets"][ds["id"]])
        datasets[ds["id"]] = result
        conflicts.extend(problems)
        events.extend(changes)
        all_views.extend(views)
    if not conflicts:
        validate_refs(cfg, datasets)
    return datasets, all_views, conflicts, events


def snapshots(root, paths):
    return {str(p.relative_to(root)): base64.b64encode(p.read_bytes()).decode("ascii") if p.exists() else None for p in paths}


def restore(root, files):
    for name, data in files.items():
        path = scoped(root, name)
        if data is None:
            if path.exists():
                path.unlink()
        else:
            atomic(path, base64.b64decode(data))


def run_checks(root, cfg):
    results = []
    for check in cfg.get("checks", []):
        argv = [sys.executable if a == "{python}" else a for a in check["argv"]]
        try:
            result = subprocess.run(argv, cwd=root, capture_output=True, text=True,
                                    encoding="utf-8", errors="replace", timeout=check.get("timeout", 120), shell=False)
        except subprocess.TimeoutExpired as exc:
            raise SyncError(f"Check timed out: {check.get('id', argv[0])}") from exc
        record = {"id": check.get("id", argv[0]), "returncode": result.returncode,
                  "stdout": result.stdout[-6000:], "stderr": result.stderr[-6000:]}
        results.append(record)
        if result.returncode:
            raise SyncError(f"Check failed: {json.dumps(record, ensure_ascii=False)}")
    return results


def execute(args):
    manifest_path = Path(args.manifest).resolve()
    if args.command == "recover":
        root, cfg, signature = manifest_path.parent, None, None
    else:
        root, cfg, signature = config(manifest_path)
    directory = root / ".design-sync"
    if not directory.resolve().is_relative_to(root):
        raise SyncError("State directory escapes project")
    directory.mkdir(exist_ok=True)
    state_path, pending_path = directory / "state.json", directory / "pending.json"
    lock = directory / "sync.lock"
    try:
        fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL)
    except FileExistsError as exc:
        raise SyncError("Sync lock exists; confirm no sync process is running before removing a stale lock") from exc
    try:
        with os.fdopen(fd, "w") as stream:
            stream.write(str(os.getpid()))
        if args.command == "recover":
            if not pending_path.exists():
                return {"status": "nothing-to-recover"}, 0
            pending = load_json(pending_path.read_bytes())
            restore(root, pending["before"])
            if pending["state_before"] is None:
                state_path.unlink(missing_ok=True)
            else:
                atomic(state_path, base64.b64decode(pending["state_before"]))
            pending_path.unlink()
            return {"status": "recovered", "note": "Reload UE assets / close and reopen the workbook after file recovery"}, 0
        if pending_path.exists():
            raise SyncError("Interrupted transaction exists; inspect pending.json and run recover first")
        if args.command in ("init", "rebaseline"):
            if args.command == "init" and state_path.exists():
                raise SyncError("Baseline exists; init never resets it")
            if args.command == "rebaseline" and not args.accept_reviewed_migration:
                raise SyncError("Rebaseline requires --accept-reviewed-migration after reviewing all existing drift")
            datasets = {}
            for ds in cfg["datasets"]:
                views = [View(root, ds, v) for v in ds["views"]]
                if any(v.records != views[0].records for v in views[1:]):
                    raise SyncError(f"Initial views disagree: {ds['id']}; reconcile explicitly, then initialize")
                datasets[ds["id"]] = views[0].records
            validate_refs(cfg, datasets)
            state = {"version": 1, "manifest_hash": signature, "datasets": datasets,
                     "view_hashes": {v["path"]: digest(scoped(root, v["path"]).read_bytes()) for d in cfg["datasets"] for v in d["views"]}}
            if args.command == "rebaseline" and state_path.exists():
                atomic(directory / f"state-before-migration-{time.time_ns()}.json", state_path.read_bytes())
            atomic(state_path, encode(state))
            return {"status": "baseline-created", "note": "Baseline creation does not execute or validate UE gameplay"}, 0
        if not state_path.exists():
            raise SyncError("No baseline; reconcile views and run init")
        state = load_json(state_path.read_bytes())
        datasets, views, conflicts, events = inspect(root, cfg, state, signature)
        report = {"status": "conflict" if conflicts else "drift" if events else "clean", "changes": events, "conflicts": conflicts}
        if conflicts:
            return report, 2
        if args.command in ("plan", "check"):
            return report, 1 if args.command == "check" and events else 0
        outputs = {v.path: v.render(datasets[v.ds["id"]]) for v in views}
        tracked = set(outputs) | {manifest_path} | {scoped(root, name) for name in cfg.get("rollback_paths", [])}
        before = snapshots(root, tracked)
        if digest((root / Path(args.manifest).name).read_bytes()) != signature:
            raise SyncError("Manifest changed during planning")
        if any(v.path.read_bytes() != v.raw for v in views):
            raise SyncError("A source changed during planning; rerun")
        pending = {"before": before, "state_before": base64.b64encode(state_path.read_bytes()).decode("ascii")}
        atomic(pending_path, encode(pending))
        try:
            for path, data in outputs.items():
                if data != path.read_bytes():
                    atomic(path, data)
            checks = run_checks(root, cfg)
            # Validators/UE importers may not silently modify the synchronized source data.
            for ds in cfg["datasets"]:
                for view_cfg in ds["views"]:
                    if View(root, ds, view_cfg).records != datasets[ds["id"]]:
                        raise SyncError("A validation command changed synchronized data")
            if digest((root / Path(args.manifest).name).read_bytes()) != signature:
                raise SyncError("A validation command changed the manifest")
            state["datasets"] = datasets
            state["view_hashes"] = {v.cfg["path"]: digest(v.path.read_bytes()) for v in views}
            atomic(state_path, encode(state))
            atomic(directory / "last-transaction.json", encode({**pending, "after": snapshots(root, tracked), "manifest_hash": signature}))
            atomic(directory / "last-report.json", encode({**report, "status": "applied", "checks": checks}))
            pending_path.unlink()
        except BaseException:
            restore(root, before)
            atomic(state_path, base64.b64decode(pending["state_before"]))
            pending_path.unlink()
            raise
        return {**report, "status": "applied", "checks": checks,
                "validation": "configured checks passed" if checks else "data only; no game checks configured"}, 0
    finally:
        lock.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("init", "plan", "check", "apply", "recover", "rebaseline"))
    parser.add_argument("--manifest", required=True, help="Manifest lives in project root; paths are relative to it")
    parser.add_argument("--accept-reviewed-migration", action="store_true")
    args = parser.parse_args()
    try:
        report, code = execute(args)
    except (SyncError, OSError, ValueError, KeyError, TypeError, ImportError) as exc:
        report, code = {"status": "error", "error": str(exc)}, 3
    print(json.dumps(report, ensure_ascii=True, indent=2, allow_nan=False))
    return code


if __name__ == "__main__":
    sys.exit(main())
